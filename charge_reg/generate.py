"""
Generate the multi-tab regularisation xlsx.
Tabs: Paramètres | Récapitulatif {year} | one tab per occupied lot
Styling matches the reference regularisation_charges_ormes_a_2025_v3.xlsx.
"""

import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Colour palette (from reference file) ──────────────────────────────────────
_NAVY   = "1B3A6B"   # dark navy  — title bg, header bg, section separators
_BLUE   = "2E5FA3"   # medium blue — category rows, provisions header, solde
_LBLUE  = "D0DDEF"   # light blue  — subtotals, subtitle rows
_LGRAY  = "F2F4F8"   # very light grey — invoice rows in récap
_GREEN  = "D5F5E3"   # light green — Occupé status
_RED    = "FDECEA"   # light red   — Vacant status

# ── Pre-built fills ────────────────────────────────────────────────────────────
_fill_navy  = PatternFill("solid", fgColor=_NAVY)
_fill_blue  = PatternFill("solid", fgColor=_BLUE)
_fill_lblue = PatternFill("solid", fgColor=_LBLUE)
_fill_lgray = PatternFill("solid", fgColor=_LGRAY)
_fill_green = PatternFill("solid", fgColor=_GREEN)
_fill_red   = PatternFill("solid", fgColor=_RED)

# ── Pre-built fonts ────────────────────────────────────────────────────────────
def _fn(size=9, bold=False, color="000000"):
    return Font(name="Calibri", size=size, bold=bold, color=color)

F_TITLE_RECAP = _fn(13, True, "FFFFFF")   # récap / paramètres main title
F_TITLE_TENANT= _fn(12, True, "FFFFFF")   # tenant tab title
F_HDRS        = _fn(8,  True, "FFFFFF")   # column headers (white on navy)
F_CAT_RECAP   = _fn(11, True, "FFFFFF")   # ── CATEGORY ── in récap
F_CAT_TENANT  = _fn(9,  True, "FFFFFF")   # category rows in tenant tab
F_SUBTOT      = _fn(9,  True, _NAVY)      # subtotals (navy text)
F_TOTAL_ROW   = _fn(10, True, "FFFFFF")   # grand total row (white on navy)
F_SOLDE       = _fn(10, True, "FFFFFF")   # solde row (white on blue)
F_NAVY_SM     = _fn(9,  False, _NAVY)     # subtitle / info text
F_POSTE_A     = _fn(9,  False, "222222")  # poste name col A
F_VAL         = _fn(9,  False, "000000")  # values / formulas
F_GRAY        = _fn(9,  False, "AAAAAA")  # placeholder ("Aucune facture")
F_BLUE_VAL    = _fn(9,  False, "000000")  # formula values (same as F_VAL)

# ── Number / date formats ──────────────────────────────────────────────────────
MONEY_FMT = '#,##0.00;(#,##0.00)'
DATE_FMT  = 'DD/MM/YYYY'
PCT_FMT   = '0.00000'

# ── Alignment shortcuts ────────────────────────────────────────────────────────
_AL_CTR   = Alignment(horizontal="center", vertical="center")
_AL_LEFT  = Alignment(horizontal="left",   vertical="center")
_AL_RIGHT = Alignment(horizontal="right",  vertical="center")
_AL_WRAP  = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ── Low-level helpers ──────────────────────────────────────────────────────────

def _c(ws, row, col, value=None, font=None, fill=None, fmt=None, align=None):
    c = ws.cell(row=row, column=col, value=value)
    if font:  c.font  = font
    if fill:  c.fill  = fill
    if fmt:   c.number_format = fmt
    if align: c.alignment = align
    return c


def _money(ws, row, col, value, font=None):
    return _c(ws, row, col, value, font=font or F_VAL, fmt=MONEY_FMT,
              align=_AL_RIGHT)


def _to_excel_date(val):
    if val is None:
        return None
    if isinstance(val, datetime.datetime):
        return val.date()
    if isinstance(val, datetime.date):
        return val
    if isinstance(val, str):
        try:
            return datetime.date.fromisoformat(val)
        except (ValueError, AttributeError):
            return None
    return None


def _set_col_widths(ws, widths: dict):
    """widths = {"A": 28, "B": 13, ...}"""
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _row_h(ws, row, height):
    ws.row_dimensions[row].height = height


# ── Formula builders ───────────────────────────────────────────────────────────

def _prorata_formula(year: int, param_row: int, is_fluid: bool) -> str:
    """Excel prorata formula. Fluids use col H (effective/MAD start), others use col F (bail start)."""
    start_col = "H" if is_fluid else "F"
    return (
        f"=MAX(0,MIN(DATE({year},12,31),"
        f"IF(Paramètres!$I${param_row}>0,Paramètres!$I${param_row},DATE({year},12,31)))"
        f"-MAX(DATE({year},1,1),"
        f"IF(Paramètres!${start_col}${param_row}>0,"
        f"Paramètres!${start_col}${param_row},DATE({year},1,1)))+1)/365"
    )


def _cle_formula(param_row: int, base_surface, site_surface: float) -> str:
    """Clé = lot_surface / base_surface.  $C$3 = site total surface in Paramètres."""
    if base_surface is None or abs(float(base_surface) - float(site_surface)) < 10:
        return f"=(Paramètres!$C${param_row}/Paramètres!$C$3)"
    return f"=(Paramètres!$C${param_row}/{round(float(base_surface), 2)})"


# ── Paramètres sheet ───────────────────────────────────────────────────────────

def build_parametres_sheet(ws, enriched: dict) -> dict:
    """
    Build the Paramètres sheet.
    Returns {lot_name: row_number} so lot tabs can reference surface/dates by row.

    Structure:
      Row 1   : title (navy)
      Row 2   : Site  (blue label | value col C)
      Row 3   : Surface totale  ← $C$3 referenced by clé formula
      Row 4   : Année
      Row 5   : (blank)
      Row 6   : column headers (navy)
      Row 7+  : one row per lot
    """
    site  = enriched["site"]
    year  = enriched.get("year", 2025)
    lots  = enriched["lots"]
    surface_totale = enriched.get("surface_totale") or sum(
        info.get("surface") or 0 for info in lots.values()
    )

    # Row 1 — title
    ws.merge_cells("A1:I1")
    _c(ws, 1, 1, f"{site} — PARAMÈTRES SITE {year}",
       font=F_TITLE_RECAP, fill=_fill_navy, align=_AL_CTR)
    _row_h(ws, 1, 27.75)

    # Rows 2-4 — site info
    for row, label, value in [
        (2, "Site",                       site),
        (3, "Surface totale site (m²)",   surface_totale),
        (4, "Année",                      year),
    ]:
        _c(ws, row, 1, label, font=_fn(9, True, "FFFFFF"), fill=_fill_blue, align=_AL_LEFT)
        _c(ws, row, 3, value, font=_fn(9, True, "000000"), align=_AL_LEFT)
        _row_h(ws, row, 15.8)

    # Row 5 — blank
    _row_h(ws, 5, 7.55)

    # Row 6 — column headers
    headers = [
        "Lot", "Statut", "Surface (m²)", "Bâtiment", "Locataire",
        "Début de bail", "Fin de bail",
        "Début effectif (MAD/bail)", "Fin effective",
    ]
    for col, h in enumerate(headers, 1):
        _c(ws, 6, col, h, font=F_HDRS, fill=_fill_navy,
           align=_AL_WRAP if col > 5 else _AL_CTR)
    _row_h(ws, 6, 19.55)

    # Rows 7+ — one row per lot
    lot_rows: dict[str, int] = {}
    for row_idx, (lot_name, info) in enumerate(lots.items(), start=7):
        lot_rows[lot_name] = row_idx

        _c(ws, row_idx, 1, lot_name, font=_fn(9, True, "000000"), align=_AL_LEFT)

        statut = info.get("statut", "Vacant")
        stat_fill = _fill_green if statut == "Occupé" else _fill_red
        _c(ws, row_idx, 2, statut, font=_fn(9, True, "000000"),
           fill=stat_fill, align=_AL_CTR)

        _c(ws, row_idx, 3, info.get("surface"), font=F_VAL, align=_AL_RIGHT)
        _c(ws, row_idx, 4, info.get("batiment") or "", font=F_VAL, align=_AL_CTR)
        _c(ws, row_idx, 5, info.get("tenant") or "", font=F_VAL, align=_AL_LEFT)

        bail_start      = _to_excel_date(info.get("bail_start"))
        bail_end        = _to_excel_date(info.get("bail_end"))
        mad_start       = _to_excel_date(info.get("mad_start"))
        effective_start = mad_start or bail_start

        for col, d in [(6, bail_start), (7, bail_end),
                       (8, effective_start), (9, bail_end)]:
            cell = ws.cell(row=row_idx, column=col, value=d)
            cell.font = F_VAL
            if d:
                cell.number_format = DATE_FMT
            cell.alignment = _AL_CTR

        _row_h(ws, row_idx, 15.05)

    # Column widths
    _set_col_widths(ws, {
        "A": 20, "B": 14, "C": 14, "D": 12, "E": 22,
        "F": 16, "G": 16, "H": 16, "I": 16,
    })

    return lot_rows


# ── Récapitulatif sheet ────────────────────────────────────────────────────────

def build_recapitulatif_sheet(ws, enriched: dict) -> dict:
    """
    Build the Récapitulatif {year} sheet with all invoices and subtotals.
    Returns {"subtotals": {(categorie, poste): subtotal_row}}
    """
    year    = enriched.get("year", 2025)
    site    = enriched["site"]
    charges = [c for c in enriched["charges"]
               if (c.get("realise_ht") or 0) > 0 or c.get("invoices")]

    # Row 1 — title
    ws.merge_cells("A1:G1")
    _c(ws, 1, 1, f"{site} — RÉCAPITULATIF DES CHARGES {year}",
       font=F_TITLE_RECAP, fill=_fill_navy, align=_AL_CTR)
    _row_h(ws, 1, 27.75)

    # Row 2 — subtitle
    ws.merge_cells("A2:G2")
    _c(ws, 2, 1,
       f"Site : {site}  |  Période : 01/01/{year} – 31/12/{year}",
       font=F_NAVY_SM, fill=_fill_lblue, align=_AL_LEFT)
    _row_h(ws, 2, 15.8)

    # Row 3 — column headers
    hdr_labels = ["Fournisseur", "Date facture", "N° facture",
                  "Description", "Montant HT (€)", "TVA (€)", "Montant TTC (€)"]
    for col, lbl in enumerate(hdr_labels, 1):
        _c(ws, 3, col, lbl, font=F_HDRS, fill=_fill_navy, align=_AL_WRAP)
    _row_h(ws, 3, 19.55)

    _set_col_widths(ws, {
        "A": 28, "B": 13, "C": 22, "D": 45, "E": 14, "F": 12, "G": 14,
    })

    current_row = 4
    current_cat = None
    poste_subtotal_rows: dict[tuple, int] = {}

    for charge in charges:
        cat     = charge.get("categorie") or ""
        poste   = charge.get("poste", "")
        invoices = charge.get("invoices", [])
        realise  = charge.get("realise_ht", 0) or 0

        if realise == 0 and not invoices:
            continue

        # Category section separator (e.g. ── FLUIDES ──)
        if cat != current_cat:
            current_cat = cat
            ws.merge_cells(f"A{current_row}:G{current_row}")
            _c(ws, current_row, 1, f"── {cat} ──",
               font=F_CAT_RECAP, fill=_fill_navy, align=_AL_CTR)
            _row_h(ws, current_row, 21.75)
            current_row += 1

        # Poste subheader
        _c(ws, current_row, 1, f"  {poste}",
           font=_fn(9, True, "FFFFFF"), fill=_fill_blue, align=_AL_LEFT)
        _row_h(ws, current_row, 18.0)
        current_row += 1

        if invoices:
            first_inv_row = current_row
            for inv in invoices:
                _c(ws, current_row, 1, inv.get("fournisseur", ""),
                   font=F_VAL, fill=_fill_lgray, align=_AL_LEFT)
                date_val = inv.get("date")
                if date_val:
                    dc = ws.cell(row=current_row, column=2, value=date_val)
                    dc.font  = F_VAL
                    dc.fill  = _fill_lgray
                    dc.number_format = DATE_FMT
                    dc.alignment = _AL_CTR
                else:
                    _c(ws, current_row, 2, None, fill=_fill_lgray)
                _c(ws, current_row, 3, inv.get("num_facture", ""),
                   font=F_VAL, fill=_fill_lgray, align=_AL_LEFT)
                _c(ws, current_row, 4, inv.get("description", ""),
                   font=F_VAL, fill=_fill_lgray, align=_AL_LEFT)
                _money(ws, current_row, 5, inv.get("realise_ht", 0)).fill = _fill_lgray
                tva = inv.get("tva") or 0
                if tva:
                    _money(ws, current_row, 6, tva).fill = _fill_lgray
                else:
                    _c(ws, current_row, 6, None, fill=_fill_lgray)
                ttc = inv.get("ttc") or 0
                if ttc:
                    _money(ws, current_row, 7, ttc).fill = _fill_lgray
                else:
                    _c(ws, current_row, 7, None, fill=_fill_lgray)
                _row_h(ws, current_row, 15.05)
                current_row += 1

            last_inv_row = current_row - 1
            # Subtotal
            st_row = current_row
            ws.merge_cells(f"A{st_row}:D{st_row}")
            _c(ws, st_row, 1, f"    Sous-total — {poste}",
               font=F_SUBTOT, fill=_fill_lblue, align=_AL_LEFT)
            for col, letter in [(5, "E"), (6, "F"), (7, "G")]:
                c = ws.cell(row=st_row, column=col)
                c.value         = f"=SUM({letter}{first_inv_row}:{letter}{last_inv_row})"
                c.font          = F_SUBTOT
                c.number_format = MONEY_FMT
                c.fill          = _fill_lblue
                c.alignment     = _AL_RIGHT
            _row_h(ws, st_row, 15.8)
            current_row += 1

        else:
            # No invoices — placeholder + static subtotal
            _c(ws, current_row, 4, "— Aucune facture —",
               font=F_GRAY, fill=_fill_lgray, align=_AL_LEFT)
            _row_h(ws, current_row, 13.5)
            current_row += 1

            st_row = current_row
            ws.merge_cells(f"A{st_row}:D{st_row}")
            _c(ws, st_row, 1, f"    Sous-total — {poste}",
               font=F_SUBTOT, fill=_fill_lblue, align=_AL_LEFT)
            c5 = ws.cell(row=st_row, column=5, value=realise)
            c5.font = F_SUBTOT; c5.number_format = MONEY_FMT
            c5.fill = _fill_lblue; c5.alignment = _AL_RIGHT
            _row_h(ws, st_row, 15.8)
            current_row += 1

        poste_subtotal_rows[(cat, poste)] = st_row

    ws.freeze_panes = "A4"
    return {"subtotals": poste_subtotal_rows}


# ── Per-lot (tenant) sheet ─────────────────────────────────────────────────────

def build_tenant_sheet(
    ws,
    lot_name: str,
    lot_info: dict,
    charges: list,
    provisions_ht: float,
    year: int,
    site: str,
    param_row: int,
    recap_subtotals: dict,
    site_surface: float,
    recap_tab_name: str,
    provisions_quarterly: dict | None = None,
):
    """Build one per-lot regularisation sheet with cross-tab formula references."""
    tenant  = lot_info.get("tenant", "LOCATAIRE")
    surface = lot_info.get("surface") or 0

    # Column widths
    _set_col_widths(ws, {"A": 34, "B": 14, "C": 15, "E": 12, "F": 15})

    # Row 1 — title
    ws.merge_cells("A1:F1")
    _c(ws, 1, 1, f"RÉGULARISATION DE CHARGES — {lot_name.upper()}",
       font=F_TITLE_TENANT, fill=_fill_navy, align=_AL_CTR)
    _row_h(ws, 1, 25.55)

    # Row 2 — static info
    ws.merge_cells("A2:F2")
    _c(ws, 2, 1,
       f"Site : {site}  |  Année {year}  |  Lot : {lot_name}  |  "
       f"Surface lot {lot_name} : {surface} m²  |  ",
       font=F_NAVY_SM, fill=_fill_lblue, align=_AL_LEFT)
    _row_h(ws, 2, 15.8)

    # Row 3 — tenant/status formula from Paramètres
    ws.merge_cells("A3:F3")
    c3 = ws.cell(row=3, column=1)
    c3.value     = (f"=Paramètres!$B${param_row}"
                    f"&IF(Paramètres!$E${param_row}<>\"\","
                    f"\" — \"&Paramètres!$E${param_row},\"\")")
    c3.font      = F_NAVY_SM
    c3.fill      = _fill_lblue
    c3.alignment = _AL_LEFT
    _row_h(ws, 3, 15.8)

    # Row 4 — column headers
    hdr_labels = ["Poste", "Clé répartition",
                  "Total site HT", "QP annuelle", "Période", "QP locataire"]
    for col, lbl in enumerate(hdr_labels, 1):
        _c(ws, 4, col, lbl, font=F_HDRS, fill=_fill_navy, align=_AL_WRAP)
    _row_h(ws, 4, 19.55)

    current_row = 5
    current_cat = None
    category_rows: dict[str, dict] = {}

    _EXCL = ("NON RECUP", "NON RÉCUP", "CAPEX")

    # ── Charge rows ────────────────────────────────────────────────────────────
    for charge in charges:
        cat      = charge.get("categorie") or ""
        poste    = charge.get("poste", "")
        realise  = charge.get("realise_ht", 0) or 0
        is_fluid = charge.get("is_fluid", False)
        base_surface = charge.get("base_surface")

        cat_upper = cat.upper()
        if any(ex in cat_upper for ex in _EXCL):
            continue
        if realise == 0 and not charge.get("invoices"):
            continue

        # Category header
        if cat != current_cat:
            current_cat = cat
            ws.merge_cells(f"A{current_row}:F{current_row}")
            _c(ws, current_row, 1, cat,
               font=F_CAT_TENANT, fill=_fill_blue, align=_AL_LEFT)
            _row_h(ws, current_row, 18.0)
            category_rows[cat] = {"header_row": current_row, "detail_rows": []}
            current_row += 1

        det_row = current_row
        category_rows[cat]["detail_rows"].append(det_row)

        # Col A — poste name
        _c(ws, det_row, 1, f"  {poste}", font=F_POSTE_A, align=_AL_LEFT)

        # Col B — Clé de répartition
        cb = ws.cell(row=det_row, column=2)
        cb.value         = _cle_formula(param_row, base_surface, site_surface)
        cb.font          = F_VAL
        cb.number_format = PCT_FMT
        cb.alignment     = _AL_CTR

        # Col C — Total site HT (references Récapitulatif subtotal)
        recap_key = (cat, poste)
        cc = ws.cell(row=det_row, column=3)
        if recap_key in recap_subtotals:
            cc.value = f"='{recap_tab_name}'!E{recap_subtotals[recap_key]}"
        else:
            cc.value = realise
        cc.font          = F_VAL
        cc.number_format = MONEY_FMT
        cc.alignment     = _AL_RIGHT

        # Col D — QP annuelle = B × C
        cd = ws.cell(row=det_row, column=4)
        cd.value = f"=B{det_row}*C{det_row}"
        cd.font  = F_VAL; cd.number_format = MONEY_FMT; cd.alignment = _AL_RIGHT

        # Col E — Prorata
        ce = ws.cell(row=det_row, column=5)
        ce.value = _prorata_formula(year, param_row, is_fluid)
        ce.font  = F_VAL; ce.number_format = "0.00000000"; ce.alignment = _AL_CTR

        # Col F — QP locataire = D × E
        cf = ws.cell(row=det_row, column=6)
        cf.value = f"=D{det_row}*E{det_row}"
        cf.font  = F_VAL; cf.number_format = MONEY_FMT; cf.alignment = _AL_RIGHT

        _row_h(ws, det_row, 15.05)
        current_row += 1

    # ── Subtotals per category ─────────────────────────────────────────────────
    subtotal_rows: list[int] = []
    current_row += 1   # blank line
    for cat, info in category_rows.items():
        if not info["detail_rows"]:
            continue
        st_row = current_row
        cat_upper = cat.upper()
        if not any(ex in cat_upper for ex in _EXCL):
            subtotal_rows.append(st_row)
        first = info["detail_rows"][0]
        last  = info["detail_rows"][-1]

        ws.merge_cells(f"A{st_row}:B{st_row}")
        _c(ws, st_row, 1, f"Sous-total — {cat}",
           font=F_SUBTOT, fill=_fill_lblue, align=_AL_LEFT)
        for col, src in [(3, "C"), (4, "D"), (6, "F")]:
            c = ws.cell(row=st_row, column=col)
            c.value         = f"=SUM({src}{first}:{src}{last})"
            c.font          = F_SUBTOT
            c.number_format = MONEY_FMT
            c.fill          = _fill_lblue
            c.alignment     = _AL_RIGHT
        _row_h(ws, st_row, 18.0)
        current_row += 1

    # ── Grand total ────────────────────────────────────────────────────────────
    current_row += 1
    total_row = current_row
    _row_h(ws, total_row - 1, 7.55)   # blank spacer

    ws.merge_cells(f"A{total_row}:B{total_row}")
    _c(ws, total_row, 1,
       "TOTAL CHARGES RÉCUPÉRABLES RÉELLES (€ HT)",
       font=F_TOTAL_ROW, fill=_fill_navy, align=_AL_LEFT)
    if subtotal_rows:
        refs_c = "+".join(f"C{r}" for r in subtotal_rows)
        refs_d = "+".join(f"D{r}" for r in subtotal_rows)
        refs_f = "+".join(f"F{r}" for r in subtotal_rows)
        for col, expr in [(3, refs_c), (4, refs_d), (6, refs_f)]:
            c = ws.cell(row=total_row, column=col)
            c.value         = f"={expr}"
            c.font          = F_TOTAL_ROW
            c.number_format = MONEY_FMT
            c.fill          = _fill_navy
            c.alignment     = _AL_RIGHT
    # Navy fill on empty cells too (cols 5 & empty cols)
    for col in [5]:
        ws.cell(row=total_row, column=col).fill = _fill_navy
    _row_h(ws, total_row, 21.75)

    # ── Provisions section ─────────────────────────────────────────────────────
    current_row += 2
    _row_h(ws, current_row - 1, 7.55)   # blank spacer

    prov_hdr = current_row
    ws.merge_cells(f"A{prov_hdr}:B{prov_hdr}")
    _c(ws, prov_hdr, 1, "Provisions pour charges versées",
       font=_fn(9, True, "FFFFFF"), fill=_fill_blue, align=_AL_LEFT)
    _c(ws, prov_hdr, 4, "Montant annuel (€)",
       font=F_HDRS, fill=_fill_blue, align=_AL_CTR)
    _c(ws, prov_hdr, 6, "Montant locataire (€)",
       font=F_HDRS, fill=_fill_blue, align=_AL_CTR)
    _row_h(ws, prov_hdr, 18.0)
    current_row += 1

    q_labels = ["  1er trimestre (T1)", "  2e trimestre (T2)",
                "  3e trimestre (T3)", "  4e trimestre (T4)"]
    q_keys   = ["Q1", "Q2", "Q3", "Q4"]
    q_rows   = []
    for i, label in enumerate(q_labels):
        qr = current_row
        q_rows.append(qr)
        ws.merge_cells(f"A{qr}:B{qr}")
        _c(ws, qr, 1, label, font=F_VAL, align=_AL_LEFT)
        # Write provision amount only to col F (QP locataire)
        if provisions_quarterly and q_keys[i] in provisions_quarterly:
            q_val = provisions_quarterly[q_keys[i]]
            if q_val:
                _money(ws, qr, 6, q_val)
        elif i == 0 and provisions_ht and not provisions_quarterly:
            _money(ws, qr, 6, provisions_ht)
        _row_h(ws, qr, 15.8)
        current_row += 1

    # Provisions total
    prov_total_row = current_row
    ws.merge_cells(f"A{prov_total_row}:C{prov_total_row}")
    _c(ws, prov_total_row, 1, "Total provisions pour charges versées →",
       font=F_SUBTOT, fill=_fill_lblue, align=_AL_LEFT)
    c6 = ws.cell(row=prov_total_row, column=6)
    c6.value         = f"=SUM(F{q_rows[0]}:F{q_rows[-1]})"
    c6.font          = F_SUBTOT
    c6.number_format = MONEY_FMT
    c6.fill          = _fill_lblue
    c6.alignment     = _AL_RIGHT
    _row_h(ws, prov_total_row, 18.0)
    current_row += 2

    # ── Solde ─────────────────────────────────────────────────────────────────
    _row_h(ws, current_row - 1, 7.55)

    solde_row = current_row
    ws.merge_cells(f"A{solde_row}:C{solde_row}")
    c_lbl = ws.cell(row=solde_row, column=1)
    c_lbl.value     = (f"=IF(F{total_row}-F{prov_total_row}>0,"
                       f"\"SOLDE DÉBITEUR\",\"SOLDE CRÉDITEUR\")")
    c_lbl.font      = F_SOLDE
    c_lbl.fill      = _fill_blue
    c_lbl.alignment = _AL_LEFT

    c_f = ws.cell(row=solde_row, column=6)
    c_f.value         = f"=F{total_row}-F{prov_total_row}"
    c_f.font          = F_SOLDE
    c_f.number_format = MONEY_FMT
    c_f.fill          = _fill_blue
    c_f.alignment     = _AL_RIGHT
    _row_h(ws, solde_row, 21.75)

    ws.freeze_panes = "A5"


# ── Top-level generate() ───────────────────────────────────────────────────────

def generate(enriched: dict, output_path: str | None = None) -> str:
    site   = enriched["site"]
    year   = enriched.get("year", 2025)
    charges = enriched["charges"]
    provisions = enriched.get("provisions", {})
    lots   = enriched["lots"]
    surface_totale = enriched.get("surface_totale") or sum(
        info.get("surface") or 0 for info in lots.values()
    )

    if output_path is None:
        out_dir = Path(enriched["filepath"]).parent
        output_path = str(out_dir / f"regularisation_{site.replace(' ', '_')}_{year}.xlsx")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    recap_tab_name = f"Récapitulatif {year}"

    ws_params = wb.create_sheet("Paramètres")
    lot_rows = build_parametres_sheet(ws_params, enriched)

    ws_recap = wb.create_sheet(recap_tab_name)
    recap_info = build_recapitulatif_sheet(ws_recap, enriched)
    recap_subtotals = recap_info["subtotals"]

    occupied_lots = {
        name: info
        for name, info in lots.items()
        if info.get("tenant") and info.get("statut", "Occupé") == "Occupé"
    }

    if not occupied_lots:
        ws_empty = wb.create_sheet("Résumé")
        ws_empty["A1"] = f"Aucun lot occupé trouvé pour {site} {year}"
    else:
        for lot_name, lot_info in occupied_lots.items():
            ws = wb.create_sheet(lot_name[:31])
            tenant       = lot_info.get("tenant", "")
            prov_ht      = provisions.get(tenant, 0.0)
            param_row    = lot_rows.get(lot_name, 7)
            prov_quarterly = enriched.get("provisions_detail", {}).get(tenant)
            build_tenant_sheet(
                ws, lot_name, lot_info, charges, prov_ht,
                year, site, param_row, recap_subtotals,
                site_surface=float(surface_totale),
                recap_tab_name=recap_tab_name,
                provisions_quarterly=prov_quarterly,
            )

    wb.save(output_path)
    return output_path
